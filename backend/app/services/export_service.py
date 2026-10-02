from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.timeutil import day_bounds_utc
from app.models import DailyStats, Expense, HotWalletTransaction, Order, Payment
from app.services.analytics_service import AnalyticsService, Period


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list[object]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([str(x) if not isinstance(x, (int, float, type(None), str)) else x for x in r])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(
            40, max(10, max(len(str(c.value or "")) for c in col) + 2)
        )


class ExportService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session

    async def export_period(self, period: Period) -> bytes:
        lo, _ = day_bounds_utc(period.start)
        _, hi = day_bounds_utc(period.end)
        analytics = AnalyticsService(self.session)
        await analytics.ensure_days(period.start, period.end)
        wb = Workbook()
        wb.remove(wb.worksheets[0])
        dash = await analytics.dashboard(period)
        cur, prev = dash["current"], dash["previous"]
        _sheet(
            wb, "Umumiy", ["Ko'rsatkich", "Joriy davr", "Oldingi davr"],
            [[k, str(cur[k]), str(prev[k])] for k in cur if not isinstance(cur[k], dict)],
        )  # fmt: skip
        days = (
            await self.session.scalars(
                select(DailyStats)
                .where(DailyStats.day >= period.start, DailyStats.day <= period.end)
                .order_by(DailyStats.day)
            )
        ).all()
        cols = [c.name for c in DailyStats.__table__.columns if c.name not in ("revenue_by_method", "updated_at")]
        _sheet(wb, "Kunlik", cols, [[getattr(d, c) for c in cols] for d in days])
        orders = (
            await self.session.scalars(
                select(Order).where(Order.created_at >= lo, Order.created_at < hi).order_by(Order.id)
            )
        ).all()
        ocols = ["public_id", "user_id", "product_type", "plan_months", "stars_amount", "recipient_username", "status", "payment_method",
                 "price_usd", "discount_usd", "cost_usd", "profit_usd", "provider", "created_at", "paid_at", "completed_at"]  # fmt: skip
        _sheet(
            wb,
            "Buyurtmalar",
            ocols,
            [
                [getattr(getattr(o, c), "value", getattr(o, c)) if getattr(o, c) is not None else None for c in ocols]
                for o in orders
            ],
        )
        pays = (
            await self.session.scalars(
                select(Payment).where(Payment.created_at >= lo, Payment.created_at < hi).order_by(Payment.id)
            )
        ).all()
        pcols = [
            "public_id",
            "user_id",
            "purpose",
            "method",
            "status",
            "amount",
            "currency",
            "amount_usd",
            "ton_tx_hash",
            "is_late",
            "created_at",
            "confirmed_at",
        ]
        _sheet(wb, "To'lovlar", pcols, [[getattr(getattr(p, c), "value", getattr(p, c)) for c in pcols] for p in pays])
        exps = (
            await self.session.scalars(
                select(Expense).where(Expense.spent_on >= period.start, Expense.spent_on <= period.end)
            )
        ).all()
        _sheet(
            wb,
            "Xarajatlar",
            ["sana", "kategoriya", "summa_usd", "izoh"],
            [[e.spent_on, e.category.value, e.amount_usd, e.note] for e in exps],
        )
        hw = (
            await self.session.scalars(
                select(HotWalletTransaction).where(
                    HotWalletTransaction.created_at >= lo, HotWalletTransaction.created_at < hi
                )
            )
        ).all()
        _sheet(
            wb,
            "HotWallet",
            ["vaqt", "yo'nalish", "tur", "ton", "usd", "tx"],
            [[h.created_at, h.direction.value, h.kind.value, h.amount_ton, h.amount_usd, h.tx_hash] for h in hw],
        )
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()
