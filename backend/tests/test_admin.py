from datetime import timedelta
from decimal import Decimal as Dc
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core.enums import AdminRole, HotWalletTxKind, PromoType
from app.core.enums import BroadcastStatus as B
from app.core.errors import Forbidden, InvalidState, NotFound, ValidationFailed
from app.core.redis import get_redis
from app.models import Admin, PremiumPlan
from app.providers.ton.chain import IncomingTx
from app.services.admin_service import AdminService
from app.services.analytics_service import AnalyticsService, resolve_period
from app.services.broadcast_service import BroadcastService
from app.services.catalog_service import CatalogService
from app.services.export_service import ExportService
from app.services.hot_wallet_service import HotWalletService
from app.services.user_service import UserService
from tests.factories import make_user
from tests.helpers import Env, FakeMessenger, plan_id, seed


@pytest.fixture
async def env(session):
    await seed(session)
    return Env(session)


async def test_roles_and_admin_management(session, env):
    await UserService(session).ensure_owner()
    owner = UserService(session)
    assert await owner.role_of(1000) == AdminRole.OWNER
    await make_user(session, 5)
    svc = AdminService(session, env.settings)
    await svc.add_admin(1000, 5, AdminRole.SUPPORT)
    assert await owner.role_of(5) == AdminRole.SUPPORT
    with pytest.raises(Forbidden):
        await svc.add_admin(1000, 5, AdminRole.OWNER)
    with pytest.raises(NotFound):
        await svc.add_admin(1000, 999, AdminRole.ADMIN)
    # a DB row claiming owner never grants owner
    await make_user(session, 6)
    session.add(Admin(user_id=6, role=AdminRole.OWNER))
    await session.flush()
    assert await owner.role_of(6) is None
    await svc.remove_admin(1000, 5)
    assert await owner.role_of(5) is None


async def test_balance_adjust_limits_and_audit(session, env):
    u = await make_user(session, 5)
    svc = AdminService(session, env.settings)
    assert await svc.adjust_balance(1000, AdminRole.OWNER, 5, Dc(500), "gift") == Dc(500)
    with pytest.raises(Forbidden):
        await svc.adjust_balance(7, AdminRole.ADMIN, 5, Dc(60), None)
    assert await svc.adjust_balance(7, AdminRole.ADMIN, 5, Dc(-20), "fix") == Dc(480)
    with pytest.raises(Forbidden):
        await svc.ban(1, 1000, None)
    await svc.ban(1000, 5, "spam")
    await session.flush()
    await session.refresh(u)
    assert u.is_banned


async def test_plan_rules_and_catalog_one_month_coming_soon(session, env):
    cat = CatalogService(session, env.settings, env.pricing, env.rates)
    data = await cat.catalog()
    months = {p["months"]: p for p in data["plans"]}
    assert months[1]["coming_soon"] and not months[1]["available"]
    assert months[3]["available"] and months[3]["price_usd"] == "13.150000" and months[12]["price_ton"] == "10.50"
    assert [p["amount"] for p in data["star_packages"]] == [50, 500]
    svc = AdminService(session, env.settings)
    one = await plan_id(session, 1)
    with pytest.raises(ValidationFailed):
        await svc.update_plan(1000, one, is_enabled=True)
    p3 = await plan_id(session, 3)
    await svc.update_plan(1000, p3, fixed_price_usd=Dc("20"), badge="🔥")
    assert (await cat.plans())[1]["price_usd"] == "20.000000"
    with pytest.raises(ValidationFailed):
        await svc.update_plan(1000, p3, id=5)


async def test_provider_price_refresh_updates_costs(session, env):
    svc = AdminService(session, env.settings)
    assert await svc.refresh_provider_prices([env.mock]) == 4
    p = await session.get(PremiumPlan, await plan_id(session, 6))
    assert p.cost_ton == Dc("5.33") and p.cost_updated_at is not None


async def test_promo_creation_rules(session, env):
    svc = AdminService(session, env.settings)
    promo = await svc.create_promo(1000, "kuz5", PromoType.PERCENT, Dc(5), max_uses=10)
    assert promo.code == "KUZ5"
    for bad in (("a b", Dc(1)), ("OK1", Dc(0)), ("BIG", Dc(95))):
        with pytest.raises(ValidationFailed):
            await svc.create_promo(1000, bad[0], PromoType.PERCENT, bad[1])


async def test_broadcast_lifecycle_blocked_flood_and_segments(session):
    users = [await make_user(session, i, language="uz" if i % 2 else "ru") for i in range(1, 8)]
    m = FakeMessenger()
    m.blocked = {3}
    svc = BroadcastService(session, m)
    b = await svc.create(1000, {"type": "text", "text": "hi"}, {"language": ["uz"]})
    assert b.total == 4 and b.status == B.DRAFT
    assert await svc.send_batch(b) == 0  # not running yet
    await svc.start(b)
    m.flood_once = True
    first = await svc.send_batch(b, 10)
    assert first == 0  # flood wait: nothing advanced
    while await svc.send_batch(b, 2):
        pass
    assert b.status == B.COMPLETED and b.sent == 3 and b.failed == 1
    assert sorted(c for c, _ in m.content_sent) == [1, 5, 7]
    await session.refresh(users[2])
    assert users[2].is_bot_blocked
    with pytest.raises(InvalidState):
        await svc.start(b)
    # blocked users are excluded from future audiences
    assert await svc.count({"language": ["uz"]}) == 3
    b2 = await svc.create(1000, {"type": "text", "text": "x"})
    await svc.start(b2)
    await svc.pause(b2)
    await svc.cancel(b2)
    assert b2.status == B.CANCELLED


async def test_excel_export_has_expected_sheets(session, env):

    u = await make_user(session, 2)
    from app.core.enums import PaymentMethod as M
    from app.core.enums import ProductType
    from app.services.checkout_service import CreateOrderRequest
    from tests.helpers import fund

    await fund(session, u, 50)
    await env.checkout.create_order(
        u, CreateOrderRequest(ProductType.PREMIUM, M.BALANCE, plan_id=await plan_id(session, 3))
    )
    await env.fulfillment().process_next()
    from app.core.timeutil import today_local

    period = resolve_period("custom", today_local() - timedelta(days=2), today_local())
    data = await ExportService(session).export_period(period)
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Umumiy", "Kunlik", "Buyurtmalar", "To'lovlar", "Xarajatlar", "HotWallet"]
    assert wb["Buyurtmalar"].max_row == 2
    assert (await AnalyticsService(session).top_buyers(today_local() - timedelta(days=2), today_local()))[0]["id"] == 2


class FakeChain:
    def __init__(self, balance):
        self.balance = balance

    async def get_balance(self, address):
        return self.balance

    async def get_incoming(self, address, after_lt=0):
        return []


class FakeSigner:
    address = "UQHot"

    def __init__(self):
        self.sent = []

    async def send(self, to, amount_nano, comment=None, body_b64=None):
        self.sent.append((to, amount_nano, comment))
        return "ab" * 32


async def test_hot_wallet_withdraw_rules(session, env):
    await env.settings.set("hot_wallet.address", "UQHot")
    signer = FakeSigner()
    hw = HotWalletService(session, env.settings, FakeChain(Dc(125)), signer, env.rates)
    assert await hw.withdrawable() == Dc("94.95")  # 125 - 30 reserve - 0.05 fee
    with pytest.raises(Forbidden):
        await hw.withdraw_to_admin(Dc(10), 5, AdminRole.ADMIN)
    with pytest.raises(ValidationFailed):
        await hw.withdraw_to_admin(Dc(95), 1000, AdminRole.OWNER)  # exceeds withdrawable
    tx = await hw.withdraw_to_admin(Dc(50), 1000, AdminRole.OWNER)
    assert tx == "ab" * 32 and signer.sent == [("UQAdminAddress", 50_000_000_000, None)]  # destination fixed by .env
    assert await hw.withdrawn_today() == Dc(50)
    await env.settings.set("hot_wallet.daily_withdraw_limit_ton", "60")
    with pytest.raises(ValidationFailed):
        await hw.withdraw_to_admin(Dc(20), 1000, AdminRole.OWNER)
    await env.settings.set("hot_wallet.frozen", True)
    with pytest.raises(InvalidState):
        await hw.withdraw_to_admin(Dc(1), 1000, AdminRole.OWNER)


async def test_auto_sweep_only_when_enabled(session, env):
    await env.settings.set("hot_wallet.address", "UQHot")
    signer = FakeSigner()
    hw = HotWalletService(session, env.settings, FakeChain(Dc(200)), signer, env.rates)
    assert await hw.auto_sweep() is None
    await env.settings.set("hot_wallet.sweep_enabled", True)
    assert await hw.auto_sweep() is not None and signer.sent[0][1] == 169_950_000_000
    assert HotWalletTxKind.SWEEP.value == "sweep"
    _ = (IncomingTx, get_redis)
